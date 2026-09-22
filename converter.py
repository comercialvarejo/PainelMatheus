#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Converte o "Relatório de Posições" exportado do rastreador (.xlsx) no JSON
que o painel consome.

USO
    python3 converter.py caminho/do/relatorio.xlsx

    # sobrescrevendo a casa e informando feriados e manutenção:
    python3 converter.py rel.xlsx \\
        --casa "Rua Tinguaçú Castanho" --cidade-casa "Arapongas" \\
        --feriado 2026-10-12=Nossa\\ Senhora\\ Aparecida \\
        --manutencao 2026-10-05:2026-10-08

O arquivo gerado fica na mesma pasta, com o nome AAAA-MM-DD_AAAA-MM-DD.json.
Depois é só acrescentar a entrada correspondente em index.json
(o script imprime o trecho pronto para colar).

FORMATO ESPERADO DA PLANILHA
    linha 1: título e período
    linha 2: cabeçalho
    linha 3+: Veículo | Data & Hora | Endereço | Cidade | Estado |
              Evento Gerador | Velocidade | Hodômetro | Condutor

REGRAS DE ANÁLISE
    - "parada" = ciclo de ignição ligada -> desligada
    - deslocamento < 1 km no ciclo  => motor ligado parado (is_real = false)
    - alerta de micro-paradas a partir de 3 no dia
    - alerta de velocidade a partir de 5 excessos no dia
    - saída tardia >= 09:30 | retorno tardio >= 20:00
    - pernoite = posição registrada mais próxima das 03:00
    - endereço alternativo = 2º local de pernoite mais frequente, se tiver
      3 noites ou mais
"""

import argparse
import collections
import datetime
import json
import re
import sys
import unicodedata
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Falta a biblioteca openpyxl. Instale com:  pip install openpyxl")

RAIZ = Path(__file__).resolve().parent

LIM_MICRO = 3          # micro-paradas no dia para virar alerta
LIM_VEL = 5            # excessos de velocidade no dia para virar alerta
DIST_MIN_REAL = 1.0    # km — abaixo disso o ciclo não é deslocamento
H_SAIDA_TARDIA = "09:30"
H_RETORNO_TARDIO = "20:00"
JANELA_NOITE_H = 6     # tolerância ao redor das 03:00 para achar o pernoite

DIAS_PT = ["Segunda-feira", "Terça-feira", "Quarta-feira", "Quinta-feira",
           "Sexta-feira", "Sábado", "Domingo"]


# ------------------------------------------------------------------ apoio

def norm(s):
    """minúsculas, sem acento, espaços colapsados — para comparar endereços"""
    s = unicodedata.normalize("NFD", (s or "").strip().lower())
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s)


def chave(addr, city):
    return f"{norm(addr)}|{norm(city)}"


def hhmm_medio(horarios):
    if not horarios:
        return None
    m = sum(int(t[:2]) * 60 + int(t[3:]) for t in horarios) // len(horarios)
    return f"{m // 60:02d}:{m % 60:02d}"


# ------------------------------------------------------------------ leitura

def ler_planilha(caminho):
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    regs = []
    placa = None
    for r in ws.iter_rows(min_row=3, values_only=True):
        if r[1] is None or not isinstance(r[1], datetime.datetime):
            continue
        placa = placa or r[0]
        regs.append({
            "dt": r[1],
            "addr": (r[2] or "").strip(),
            "city": (r[3] or "").strip(),
            "ev": (r[5] or "").strip(),
            "hod": float(r[7] or 0),
            "cond": (r[8] or "").strip(),
        })
    if not regs:
        sys.exit("Nenhum registro de posição encontrado na planilha.")
    regs.sort(key=lambda x: x["dt"])
    return regs, placa


def montar_viagens(regs):
    """agrupa os registros em ciclos de ignição"""
    ciclos, atual = [], None
    for p in regs:
        if p["ev"] == "Ignição ligada":
            if atual:
                ciclos.append(atual)      # ligada sem desligada anterior
            atual = {"ini": p, "pts": [p], "fim": None}
        elif atual is not None:
            atual["pts"].append(p)
            if p["ev"] == "Ignição desligada":
                atual["fim"] = p
                ciclos.append(atual)
                atual = None
    if atual:
        ciclos.append(atual)

    viagens = []
    for c in ciclos:
        fim = c["fim"] or c["pts"][-1]
        dist = round(fim["hod"] - c["ini"]["hod"], 1)
        viagens.append({
            "ini_dt": c["ini"]["dt"], "fim_dt": fim["dt"],
            "start_addr": c["ini"]["addr"], "start_city": c["ini"]["city"],
            "end_addr": fim["addr"], "end_city": fim["city"],
            "dist_km": max(dist, 0.0),
            "dur_min": round(max((fim["dt"] - c["ini"]["dt"]).total_seconds() / 60, 0.0), 1),
            "is_real": dist >= DIST_MIN_REAL,
            "n_exc": sum(1 for p in c["pts"] if p["ev"] == "Velocidade máxima excedida"),
            "n_fre": sum(1 for p in c["pts"] if p["ev"] == "Frenagem Brusca"),
            "n_ace": sum(1 for p in c["pts"] if p["ev"] == "Aceleração Brusca"),
        })
    return viagens


# ------------------------------------------------------------------ análise

def converter(args):
    regs, placa = ler_planilha(args.xlsx)
    placa = args.placa or placa
    viagens = montar_viagens(regs)

    condutores = collections.Counter(p["cond"] for p in regs if p["cond"])
    condutor = args.motorista or (condutores.most_common(1)[0][0].title()
                                  if condutores else "Não informado")

    # grafia mais frequente de cada endereço (o relatório varia a acentuação)
    grafias = collections.defaultdict(collections.Counter)
    for p in regs:
        if p["addr"]:
            grafias[chave(p["addr"], p["city"])][p["addr"]] += 1

    def bonito(k):
        return grafias[k].most_common(1)[0][0] if k in grafias else "?"

    # ---------------------------------------------------------- pernoites
    def pos_noite(dia):
        alvo = datetime.datetime.combine(dia, datetime.time(3, 0))
        cands = [p for p in regs
                 if abs((p["dt"] - alvo).total_seconds()) <= JANELA_NOITE_H * 3600
                 and p["addr"]]
        return min(cands, key=lambda p: abs((p["dt"] - alvo).total_seconds())) if cands else None

    d0, d1 = regs[0]["dt"].date(), regs[-1]["dt"].date()
    datas, d = [], d0
    while d <= d1:
        datas.append(d)
        d += datetime.timedelta(days=1)

    pernoite = {}
    for d in datas:
        p = pos_noite(d)
        pernoite[d] = chave(p["addr"], p["city"]) if p else None

    cont = collections.Counter(v for v in pernoite.values() if v)
    if not cont:
        sys.exit("Não foi possível determinar nenhum pernoite.")

    # casa: informada pelo usuário ou o local de pernoite mais frequente
    if args.casa:
        casa = chave(args.casa, args.cidade_casa or regs[0]["city"])
    else:
        casa = cont.most_common(1)[0][0]

    vizinhanca = {chave(a, args.cidade_casa or regs[0]["city"]) for a in (args.vizinhanca or [])}

    # endereços onde o veículo "pernoitou" mas o hodômetro não mexeu no dia
    # inteiro são ruído de geocodificação: a mesma vaga lida como rua vizinha
    for k in list(cont):
        if k == casa or k in vizinhanca:
            continue
        noites = [d for d, v in pernoite.items() if v == k]
        moveu = any(any(v["dist_km"] > 0 for v in viagens if v["ini_dt"].date() == n)
                    for n in noites)
        if not moveu:
            vizinhanca.add(k)

    alt = next((k for k, v in cont.most_common()
                if k != casa and k not in vizinhanca and v >= 3), None)

    def eh_casa(k):
        return k == casa or k in vizinhanca

    def rotulo(k):
        if k is None:
            return "Desconhecido"
        if k == casa:
            return f"Casa ({bonito(k)})"
        if k in vizinhanca:
            return f"Casa (vizinhança - {bonito(k)}, hodômetro imóvel)"
        if alt and k == alt:
            return f"Endereço alternativo ({bonito(k)})"
        return f"Outro local: {bonito(k)}"

    # ------------------------------------------------------------- por dia
    ultimo = regs[-1]["dt"]
    dia_parcial = ultimo.date() if ultimo.hour < 18 else None
    manutencao = set(args.dias_manutencao or [])

    dias = []
    for d in datas:
        vs = [v for v in viagens if v["ini_dt"].date() == d]
        reais = [v for v in vs if v["is_real"]]
        micro = [v for v in vs if not v["is_real"]]
        exc = sum(v["n_exc"] for v in vs)
        fds = d.weekday() >= 5
        feriado = d in args.feriados
        parcial = d == dia_parcial
        manut = d in manutencao

        saida = reais[0]["ini_dt"].strftime("%H:%M") if reais else None
        retorno = reais[-1]["fim_dt"].strftime("%H:%M") if reais else None
        pk = pernoite.get(d)

        flags = []
        if manut:
            flags.append("Veículo em manutenção - dia desconsiderado como desvio")
        elif feriado:
            flags.append(f"Feriado ({args.feriados[d]}) - sem expediente previsto")
        elif not fds:
            if not vs:
                flags.append("Nenhum deslocamento registrado (dia útil sem movimentação)")
            elif not reais:
                flags.append("Nenhuma saída com deslocamento real (>1km); só partidas no local")
            if saida and saida >= H_SAIDA_TARDIA:
                flags.append(f"Saída tardia para o trabalho ({saida})")
            if retorno and retorno >= H_RETORNO_TARDIO and not parcial:
                flags.append(f"Retorno tardio ({retorno})")
        if parcial:
            flags.append(f"Dia parcial - o relatório termina às {ultimo:%H:%M}")
        if pk and not eh_casa(pk) and (alt is None or pk != alt) and not manut:
            fora = "dos 2 endereços habituais" if alt else "do endereço habitual"
            flags.append(f"Pernoite em local incomum (fora {fora})")
        if not feriado and not manut and len(micro) >= LIM_MICRO:
            flags.append(f"{len(micro)} partidas/paradas sem deslocamento relevante (<1km)")
        if exc >= LIM_VEL:
            flags.append(f"{exc} excessos de velocidade no dia")

        dias.append({
            "date": d.isoformat(), "weekday": DIAS_PT[d.weekday()], "is_weekend": fds,
            "overnight_prev_night": ("Manutenção/Oficina (" + bonito(pk) + ")"
                                     if manut and pk else rotulo(pk)),
            "first_departure_time": saida,
            "first_departure_addr": (f"{reais[0]['start_addr']}, {reais[0]['start_city']}"
                                     if reais else None),
            "return_time": retorno,
            "return_addr": (f"{reais[-1]['end_addr']}, {reais[-1]['end_city']}"
                            if reais else None),
            "return_is_home": (eh_casa(chave(reais[-1]["end_addr"], reais[-1]["end_city"]))
                               if reais else None),
            "num_stops": len(vs), "num_micro_stops": len(micro),
            "total_km": round(sum(v["dist_km"] for v in vs), 1),
            "speed_events": exc,
            "brake_events": sum(v["n_fre"] for v in vs),
            "accel_events": sum(v["n_ace"] for v in vs),
            "flags": flags,
            "in_maintenance": manut,
            "trips": [{"start": v["ini_dt"].strftime("%H:%M"),
                       "end": v["fim_dt"].strftime("%H:%M"),
                       "start_addr": v["start_addr"], "start_city": v["start_city"],
                       "end_addr": v["end_addr"], "end_city": v["end_city"],
                       "dist_km": v["dist_km"], "dur_min": v["dur_min"],
                       "is_real": v["is_real"]} for v in vs],
        })

    # -------------------------------------------------------- motor ocioso
    ocioso = sorted(
        [{"date": v["ini_dt"].date().isoformat(),
          "start": v["ini_dt"].strftime("%H:%M"), "end": v["fim_dt"].strftime("%H:%M"),
          "addr": v["start_addr"], "city": v["start_city"],
          "dist_km": v["dist_km"], "dur_min": v["dur_min"]}
         for v in viagens if not v["is_real"]],
        key=lambda e: -e["dur_min"])

    # -------------------------------------------------- episódios no alt.
    episodios = []
    if alt:
        dentro, chegada = False, None
        for p in regs:
            k = chave(p["addr"], p["city"])
            if k == alt and not dentro:
                dentro, chegada = True, p["dt"]
            elif k != alt and dentro and p["addr"]:
                horas = round((p["dt"] - chegada).total_seconds() / 3600, 1)
                noites = len({chegada.date(), p["dt"].date()}) > 1
                fds = chegada.weekday() >= 4 and horas >= 24
                episodios.append({
                    "arrival": chegada.strftime("%Y-%m-%d %H:%M"),
                    "departure": p["dt"].strftime("%Y-%m-%d %H:%M"),
                    "dwell_h": horas,
                    "kind": ("Pernoite de fim de semana / múltiplos dias" if fds
                             else "Pernoite" if noites else "Passagem rápida"),
                })
                dentro = False

    # ------------------------------------------------------------- resumo
    feriados_iso = {k.isoformat() for k in args.feriados}
    manut_iso = {k.isoformat() for k in manutencao}
    base = [x for x in dias if not x["is_weekend"]
            and x["date"] not in feriados_iso and x["date"] not in manut_iso
            and (dia_parcial is None or x["date"] != dia_parcial.isoformat())]
    com_mov = [x for x in base if x["num_stops"] > 0]

    noites = collections.Counter()
    for d, k in pernoite.items():
        if d in manutencao:
            noites["manut"] += 1
        elif k is None:
            noites["?"] += 1
        elif eh_casa(k):
            noites["casa"] += 1
        elif alt and k == alt:
            noites["alt"] += 1
        else:
            noites["outro"] += 1

    dias_alt = len({e["arrival"][:10] for e in episodios}) if episodios else 0

    summary = {
        "total_days": len(dias),
        "business_days": len(base),
        "days_no_movement": sum(1 for x in base if x["num_stops"] == 0),
        "avg_departure": hhmm_medio([x["first_departure_time"] for x in com_mov
                                     if x["first_departure_time"]]),
        "avg_return": hhmm_medio([x["return_time"] for x in com_mov if x["return_time"]]),
        "avg_stops_per_day": round(sum(x["num_stops"] for x in com_mov) / max(len(com_mov), 1), 1),
        "avg_km_per_day": round(sum(x["total_km"] for x in com_mov) / max(len(com_mov), 1), 1),
        "total_km": round(sum(x["total_km"] for x in dias), 1),
        "nights_home": noites["casa"], "nights_alt": noites["alt"],
        "nights_other": noites["outro"], "nights_maintenance": noites["manut"],
        "nights_unknown": noites["?"],
        "total_speed_events": sum(x["speed_events"] for x in dias),
        # feriado, dia parcial e manutenção explicam o dia, não são desvio
        "days_with_flags": sum(1 for x in dias if any(
            not f.startswith(("Feriado", "Dia parcial", "Veículo em manutenção"))
            for f in x["flags"])),
        "idle_total_events": len(ocioso),
        "idle_avg_min": round(sum(e["dur_min"] for e in ocioso) / max(len(ocioso), 1), 1),
        "idle_over_15min": sum(1 for e in ocioso if e["dur_min"] >= 15),
        "pm_distinct_days": dias_alt,
    }

    rel_id = f"{d0.isoformat()}_{d1.isoformat()}"
    rel = {
        "meta": {
            "id": rel_id,
            "titulo": "Relatório de Rastreamento Veicular",
            "motorista": condutor,
            "placa": placa,
            "periodo_inicio": d0.isoformat(),
            "periodo_fim": d1.isoformat(),
            "endereco_casa": f"{bonito(casa)}"
                             + (f", {args.cidade_casa}" if args.cidade_casa else ""),
            "endereco_alternativo": bonito(alt) if alt else None,
            "gerado_em": datetime.date.today().isoformat(),
        },
        "summary": summary,
        "days": dias,
        "idle_events": ocioso,
        "geo_table": [],        # exige geolocalização — preencher à mão se quiser
        "pm_episodes": episodios,
    }

    destino = RAIZ / f"{rel_id}.json"
    destino.write_text(json.dumps(rel, ensure_ascii=False, separators=(",", ":")),
                       encoding="utf-8")

    # ------------------------------------------------------------ resultado
    print(f"\nGerado: {destino.name}  ({destino.stat().st_size // 1024} KB)")
    print(f"  veículo {placa} · condutor {condutor}")
    print(f"  período {d0:%d/%m/%Y} a {d1:%d/%m/%Y}")
    print(f"  {len(dias)} dias · {len(base)} dias úteis cheios · {len(viagens)} paradas")
    print(f"  casa: {bonito(casa)}")
    if vizinhanca:
        print("  tratados como vizinhança da casa (hodômetro imóvel):")
        for k in sorted(vizinhanca):
            if k in grafias:
                print(f"    - {bonito(k)}")
    print(f"  endereço alternativo: {bonito(alt) if alt else 'nenhum detectado'}")
    print(f"  pernoites: casa={noites['casa']} alt={noites['alt']} "
          f"outro={noites['outro']} manutenção={noites['manut']}")
    print(f"  dias com alerta: {summary['days_with_flags']}")

    print("\nAcrescente este trecho em index.json, dentro de \"relatorios\":\n")
    print(json.dumps({
        "id": rel_id,
        "rotulo": f"{d0:%d/%m/%Y} – {d1:%d/%m/%Y}",
        "motorista": condutor,
        "placa": placa,
        "arquivo": destino.name,
        "inicio": d0.isoformat(),
        "fim": d1.isoformat(),
        "padrao": True,
    }, ensure_ascii=False, indent=2))
    print('\nLembre de deixar "padrao": false nos relatórios antigos.')
    print("Depois rode:  python3 validar.py")


# ------------------------------------------------------------------- CLI

def data_iso(s):
    try:
        return datetime.date.fromisoformat(s)
    except ValueError:
        raise argparse.ArgumentTypeError(f"data inválida: {s} (use AAAA-MM-DD)")


def main():
    ap = argparse.ArgumentParser(
        description="Converte o Relatório de Posições (.xlsx) no JSON do painel.",
        formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx", help="planilha exportada do rastreador")
    ap.add_argument("--casa", help="endereço de referência (casa). "
                                   "Sem isso, usa o local de pernoite mais frequente.")
    ap.add_argument("--cidade-casa", help="cidade do endereço de referência")
    ap.add_argument("--vizinhanca", nargs="*", metavar="ENDEREÇO",
                    help="ruas vizinhas que o rastreador confunde com a casa")
    ap.add_argument("--motorista", help="nome do motorista (padrão: o condutor da planilha)")
    ap.add_argument("--placa", help="placa (padrão: a da planilha)")
    ap.add_argument("--feriado", action="append", default=[], metavar="AAAA-MM-DD=Nome",
                    help="feriado a desconsiderar; pode repetir")
    ap.add_argument("--manutencao", action="append", default=[], metavar="INÍCIO:FIM",
                    help="período de manutenção (AAAA-MM-DD:AAAA-MM-DD); pode repetir")
    args = ap.parse_args()

    args.feriados = {}
    for f in args.feriado:
        if "=" not in f:
            sys.exit(f"--feriado precisa do formato AAAA-MM-DD=Nome (recebi: {f})")
        dia, nome = f.split("=", 1)
        args.feriados[data_iso(dia)] = nome

    args.dias_manutencao = []
    for m in args.manutencao:
        ini, _, fim = m.partition(":")
        ini, fim = data_iso(ini), data_iso(fim or ini)
        d = ini
        while d <= fim:
            args.dias_manutencao.append(d)
            d += datetime.timedelta(days=1)

    if not Path(args.xlsx).exists():
        sys.exit(f"Arquivo não encontrado: {args.xlsx}")

    converter(args)


if __name__ == "__main__":
    main()
